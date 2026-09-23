"""Behavior tests for the uploadable data-pptx Skill case."""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "generate.py"
SAMPLE = ROOT / "resources" / "example.json"
module_spec = importlib.util.spec_from_file_location("data_pptx_generate", SCRIPT)
generator = importlib.util.module_from_spec(module_spec)
module_spec.loader.exec_module(generator)


class DataPptxCaseTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.output = Path(self.directory.name) / "report.pptx"
        self.data = json.loads(SAMPLE.read_text(encoding="utf-8"))

    def test_sample_cli_builds_editable_charts_from_data(self) -> None:
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(SAMPLE), str(self.output)],
            check=True, capture_output=True, text=True,
        )
        status = json.loads(result.stdout)
        self.assertEqual(status["slides"], 5)
        self.assertEqual(status["charts"], 3)
        deck = Presentation(self.output)
        charts = [shape.chart for slide in deck.slides for shape in slide.shapes
                  if shape.has_chart]
        self.assertEqual(len(charts), 3)
        self.assertEqual(
            [chart.chart_type for chart in charts],
            [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.LINE_MARKERS,
             XL_CHART_TYPE.PIE],
        )
        self.assertEqual(list(charts[0].series[0].values), [390, 450, 520])
        self.assertEqual(charts[0].series[0].format.fill.fore_color.rgb,
                         RGBColor(0x32, 0x78, 0xBD))
        self.assertTrue(charts[2].plots[0].data_labels.show_percentage)
        self.assertFalse(charts[2].plots[0].data_labels.show_value)

    def test_chart_options_and_template_can_be_changed(self) -> None:
        template = Path(self.directory.name) / "company-template.pptx"
        company = Presentation()
        company.slide_width = Inches(10)
        company.slide_height = Inches(7.5)
        company.save(template)
        report = self.data
        report["template"] = str(template)
        report["slides"] = [report["slides"][1]]
        chart = report["slides"][0]["chart"]
        chart["type"] = "bar"
        chart["series"][0]["color"] = "#6B5B95"
        chart["show_legend"] = False
        chart["show_data_labels"] = False
        status = generator.build_presentation(report, self.output)
        self.assertEqual(status["charts"], 1)
        deck = Presentation(self.output)
        self.assertEqual(deck.slide_width, Inches(10))
        self.assertEqual(deck.slide_height, Inches(7.5))
        native_chart = next(shape.chart for shape in deck.slides[1].shapes
                            if shape.has_chart)
        self.assertEqual(native_chart.chart_type, XL_CHART_TYPE.BAR_CLUSTERED)
        self.assertEqual(native_chart.series[0].format.fill.fore_color.rgb,
                         RGBColor(0x6B, 0x5B, 0x95))
        self.assertFalse(native_chart.has_legend)
        self.assertFalse(native_chart.plots[0].has_data_labels)

    def test_invalid_data_and_template_are_rejected(self) -> None:
        self.data["slides"][1]["chart"]["series"][0]["values"] = [1]
        with self.assertRaisesRegex(generator.PresentationSpecError,
                                    "match categories"):
            generator.build_presentation(self.data, self.output)
        self.assertFalse(self.output.exists())

        self.data = json.loads(SAMPLE.read_text(encoding="utf-8"))
        template = Path(self.directory.name) / "nonempty.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0])
        deck.save(template)
        with self.assertRaisesRegex(generator.PresentationSpecError,
                                    "no slides"):
            generator.build_presentation(self.data, self.output, template)

    def test_upload_package_has_skill_paths_and_fits_limits(self) -> None:
        package = ROOT / "data-pptx-skill.zip"
        self.assertLess(package.stat().st_size, 4 * 1024 * 1024)
        with ZipFile(package) as archive:
            self.assertEqual(
                set(archive.namelist()),
                {"SKILL.md", "scripts/generate.py", "resources/example.json",
                 "resources/report-template.pptx"},
            )
            self.assertLess(sum(item.file_size for item in archive.infolist()),
                            4 * 1024 * 1024)
            for name in archive.namelist():
                self.assertEqual(archive.read(name), (ROOT / name).read_bytes())


if __name__ == "__main__":
    unittest.main()
