"""Generate an editable PowerPoint report from a JSON specification.

The Agent.Matrix Skill runner mounts this file under /input and collects files
written to /output. Python dependencies come from the Runner image.
"""

from __future__ import annotations

import json
import math
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


SKILL_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_ROOT / "resources" / "report-template.pptx"
DEFAULT_OUTPUT = Path("/output/report.pptx")
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}
LEGEND_POSITIONS = {
    "bottom": XL_LEGEND_POSITION.BOTTOM,
    "top": XL_LEGEND_POSITION.TOP,
    "left": XL_LEGEND_POSITION.LEFT,
    "right": XL_LEGEND_POSITION.RIGHT,
}
DEFAULT_PALETTE = ["#3278BD", "#E7A23B", "#54A582", "#9A72B0", "#D56968"]


class PresentationSpecError(ValueError):
    """Invalid report data or template."""


def color(value: str) -> RGBColor:
    if not isinstance(value, str) or len(value) != 7 or value[0] != "#":
        raise PresentationSpecError("Colors must use #RRGGBB.")
    try:
        return RGBColor.from_string(value[1:])
    except ValueError as exc:
        raise PresentationSpecError("Colors must use #RRGGBB.") from exc


def _required_text(value: Any, name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise PresentationSpecError(f"{name} must be non-empty text.")
    return value.strip()


def _position(layouts: Any, value: Any, default: int) -> Any:
    if value is None:
        value = default
    if isinstance(value, int) and not isinstance(value, bool):
        if 0 <= value < len(layouts):
            return layouts[value]
    elif isinstance(value, str):
        for layout in layouts:
            if layout.name == value:
                return layout
    raise PresentationSpecError(f"Unknown template layout: {value!r}.")


def _add_text(slide: Any, text: str, left: int, top: int, width: int,
              height: int, size: int, rgb: RGBColor, bold: bool = False,
              align: Any = PP_ALIGN.LEFT) -> Any:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    paragraph.font.name = "Noto Sans CJK SC"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb
    return box


def _add_heading(slide: Any, title: str, presentation: Any, accent: RGBColor) -> None:
    width, height = presentation.slide_width, presentation.slide_height
    _add_text(slide, title, int(width * .065), int(height * .075),
              int(width * .87), int(height * .16), 30, accent, bold=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(width * .065), int(height * .23),
                                  int(width * .15), int(height * .008))
    rule.fill.solid()
    rule.fill.fore_color.rgb = accent
    rule.line.fill.background()


def _chart_series(raw: Any, categories: list[str], chart_type: str) -> list[dict[str, Any]]:
    if not isinstance(raw, list) or not 1 <= len(raw) <= 8:
        raise PresentationSpecError("A chart needs 1 to 8 series.")
    if chart_type == "pie" and len(raw) != 1:
        raise PresentationSpecError("A pie chart needs exactly one series.")
    series = []
    for item in raw:
        if not isinstance(item, dict):
            raise PresentationSpecError("Each series must be an object.")
        name = _required_text(item.get("name"), "series.name")
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(categories):
            raise PresentationSpecError("Series values must match categories.")
        if any(isinstance(value, bool) or not isinstance(value, (int, float))
               or not math.isfinite(value) for value in values):
            raise PresentationSpecError("Series values must be finite numbers.")
        if chart_type == "pie" and any(value < 0 for value in values):
            raise PresentationSpecError("Pie values cannot be negative.")
        if "color" in item:
            color(item["color"])
        series.append({"name": name, "values": values, "color": item.get("color")})
    return series


def _add_chart(slide: Any, spec: dict[str, Any], presentation: Any,
               accent: RGBColor, palette: list[str]) -> None:
    chart_spec = spec.get("chart")
    if not isinstance(chart_spec, dict):
        raise PresentationSpecError("chart slide requires a chart object.")
    chart_type = chart_spec.get("type", "column")
    if chart_type not in CHART_TYPES:
        raise PresentationSpecError(f"Unsupported chart type: {chart_type!r}.")
    categories = chart_spec.get("categories")
    if not isinstance(categories, list) or not 1 <= len(categories) <= 50 \
            or any(not isinstance(item, str) or not item.strip() for item in categories):
        raise PresentationSpecError("Chart categories need 1 to 50 non-empty labels.")
    series = _chart_series(chart_spec.get("series"), categories, chart_type)
    data = CategoryChartData()
    data.categories = categories
    for item in series:
        data.add_series(item["name"], item["values"])

    width, height = presentation.slide_width, presentation.slide_height
    frame = slide.shapes.add_chart(
        CHART_TYPES[chart_type], int(width * .07), int(height * .28),
        int(width * .86), int(height * .59), data)
    chart = frame.chart
    chart.has_title = False
    chart.has_legend = bool(chart_spec.get("show_legend", len(series) > 1))
    if chart.has_legend:
        legend_position = chart_spec.get("legend_position", "bottom")
        if legend_position not in LEGEND_POSITIONS:
            raise PresentationSpecError("legend_position must be top, bottom, left or right.")
        chart.legend.position = LEGEND_POSITIONS[legend_position]
        chart.legend.include_in_layout = False

    show_labels = chart_spec.get("show_data_labels", False)
    if not isinstance(show_labels, bool):
        raise PresentationSpecError("show_data_labels must be boolean.")
    plot = chart.plots[0]
    plot.has_data_labels = show_labels
    if show_labels:
        plot.data_labels.show_value = chart_type != "pie"
        plot.data_labels.show_percentage = chart_type == "pie"
        plot.data_labels.font.size = Pt(11)
        if "number_format" in chart_spec and chart_type != "pie":
            plot.data_labels.number_format = _required_text(
                chart_spec["number_format"], "number_format")

    if chart_type == "pie":
        for index, point in enumerate(chart.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = color(palette[index % len(palette)])
    else:
        for index, item in enumerate(series):
            series_color = color(item["color"] or palette[index % len(palette)])
            if chart_type == "line":
                chart.series[index].format.line.color.rgb = series_color
                chart.series[index].format.line.width = Pt(3)
            else:
                chart.series[index].format.fill.solid()
                chart.series[index].format.fill.fore_color.rgb = series_color
        for key, attribute in (("axis_min", "minimum_scale"),
                               ("axis_max", "maximum_scale")):
            if key in chart_spec:
                value = chart_spec[key]
                if isinstance(value, bool) or not isinstance(value, (int, float)) \
                        or not math.isfinite(value):
                    raise PresentationSpecError(f"{key} must be a finite number.")
                setattr(chart.value_axis, attribute, value)
        if "number_format" in chart_spec:
            chart.value_axis.tick_labels.number_format = _required_text(
                chart_spec["number_format"], "number_format")

    _add_heading(slide, _required_text(spec.get("title"), "slide.title"),
                 presentation, accent)
    if spec.get("note"):
        _add_text(slide, _required_text(spec["note"], "slide.note"),
                  int(width * .07), int(height * .9), int(width * .86),
                  int(height * .07), 10, color("#667085"))


def _add_bullets(slide: Any, spec: dict[str, Any], presentation: Any,
                 accent: RGBColor) -> None:
    items = spec.get("items")
    if not isinstance(items, list) or not 1 <= len(items) <= 8:
        raise PresentationSpecError("Bullet slides need 1 to 8 items.")
    width, height = presentation.slide_width, presentation.slide_height
    _add_heading(slide, _required_text(spec.get("title"), "slide.title"),
                 presentation, accent)
    for index, item in enumerate(items):
        text = _required_text(item, "slide.items")
        top = int(height * (.32 + index * .075))
        _add_text(slide, "• " + text, int(width * .09), top,
                  int(width * .82), int(height * .068), 20, color("#263238"))


def build_presentation(spec: dict[str, Any], output: Path,
                       template_override: Path | None = None) -> dict[str, Any]:
    if not isinstance(spec, dict):
        raise PresentationSpecError("The report specification must be an object.")
    title = _required_text(spec.get("title"), "title")
    template_value = spec.get("template")
    if template_override is not None:
        template = template_override
    elif template_value is None:
        template = DEFAULT_TEMPLATE
    else:
        template = Path(_required_text(template_value, "template"))
        if not template.is_absolute():
            template = SKILL_ROOT / template
    if not template.is_file() or template.suffix.lower() != ".pptx":
        raise PresentationSpecError("template must point to an existing .pptx file.")
    try:
        presentation = Presentation(str(template))
    except Exception as exc:
        raise PresentationSpecError("Unable to open the PPTX template.") from exc
    if len(presentation.slides):
        raise PresentationSpecError("The template must contain layouts but no slides.")
    if len(presentation.slide_layouts) < 1:
        raise PresentationSpecError("The template has no slide layouts.")

    theme = spec.get("theme", {})
    if not isinstance(theme, dict):
        raise PresentationSpecError("theme must be an object.")
    accent = color(theme.get("accent", "#215B87"))
    palette = theme.get("palette", DEFAULT_PALETTE)
    if not isinstance(palette, list) or not 1 <= len(palette) <= 12:
        raise PresentationSpecError("theme.palette needs 1 to 12 colors.")
    for item in palette:
        color(item)
    layouts = spec.get("layouts", {})
    if not isinstance(layouts, dict):
        raise PresentationSpecError("layouts must be an object.")
    title_layout = _position(presentation.slide_layouts, layouts.get("title"), 0)
    content_layout = _position(presentation.slide_layouts, layouts.get("content"),
                               min(6, len(presentation.slide_layouts) - 1))

    slides = spec.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 30:
        raise PresentationSpecError("slides needs 1 to 30 content slides.")
    cover = presentation.slides.add_slide(title_layout)
    if cover.shapes.title is not None:
        cover.shapes.title.text = title
    else:
        _add_text(cover, title, int(presentation.slide_width * .08),
                  int(presentation.slide_height * .35),
                  int(presentation.slide_width * .84),
                  int(presentation.slide_height * .18), 38, accent, bold=True)
    subtitle = spec.get("subtitle")
    if subtitle is not None:
        subtitle = _required_text(subtitle, "subtitle")
        subtitle_placeholder = next((shape for shape in cover.placeholders
                                     if shape.placeholder_format.idx == 1), None)
        if subtitle_placeholder is not None:
            subtitle_placeholder.text = subtitle
        else:
            _add_text(cover, subtitle, int(presentation.slide_width * .08),
                      int(presentation.slide_height * .57),
                      int(presentation.slide_width * .84),
                      int(presentation.slide_height * .12), 20, accent)

    chart_count = 0
    for item in slides:
        if not isinstance(item, dict):
            raise PresentationSpecError("Each slide must be an object.")
        slide = presentation.slides.add_slide(content_layout)
        # Content layouts can carry inherited placeholders; clear text so the
        # chart and our report title are the only foreground content.
        for placeholder in slide.placeholders:
            if placeholder.has_text_frame:
                placeholder.text = ""
        kind = item.get("kind")
        if kind == "bullets":
            _add_bullets(slide, item, presentation, accent)
        elif kind == "chart":
            _add_chart(slide, item, presentation, accent, palette)
            chart_count += 1
        else:
            raise PresentationSpecError(f"Unsupported slide kind: {kind!r}.")

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    if output.stat().st_size > 10 * 1024 * 1024:
        output.unlink()
        raise PresentationSpecError("The PPTX exceeds Runner's 10 MiB output limit.")
    return {"output": str(output), "slides": len(presentation.slides),
            "charts": chart_count, "template": str(template)}


def main() -> int:
    if len(sys.argv) not in (2, 3):
        raise PresentationSpecError(
            "Pass one JSON object or a JSON file path; optional second argument is the output path.")
    source = sys.argv[1]
    if source.lstrip().startswith("{"):
        spec = json.loads(source)
    else:
        spec = json.loads(Path(source).read_text(encoding="utf-8"))
    output = Path(sys.argv[2]) if len(sys.argv) == 3 else DEFAULT_OUTPUT
    print(json.dumps(build_presentation(spec, output), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (PresentationSpecError, OSError, json.JSONDecodeError) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(2) from exc
